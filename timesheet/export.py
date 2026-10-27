"""
ماژول تولید فایل خروجی اکسل برای گزارش‌های سامانه ساعت کاری
"""
import io
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
import jdatetime
from django.http import HttpResponse


def export_daily_report_to_excel(daily_report, total_hours, start=None, end=None, project=None):
    """
    تولید فایل اکسل زیبا و راست‌چین از سوابق گزارش فیلترشده کارکرد
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "گزارش کارکرد"

    # فعال‌سازی نمایش راست‌چین برای زبان فارسی
    ws.views.sheetView[0].rightToLeft = True

    # تعاریف استایل و فونت‌ها
    font_family = "Tahoma"
    title_font = Font(name=font_family, size=14, bold=True, color="1E293B")
    sub_font = Font(name=font_family, size=10, italic=True, color="64748B")
    header_font = Font(name=font_family, size=10, bold=True, color="FFFFFF")
    cell_font = Font(name=font_family, size=10, color="1E293B")
    total_font = Font(name=font_family, size=11, bold=True, color="065F46")

    header_fill = PatternFill(start_color="059669", end_color="059669", fill_type="solid")
    alt_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    white_fill = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
    total_fill = PatternFill(start_color="D1FAE5", end_color="D1FAE5", fill_type="solid")

    thin_border = Border(
        left=Side(style="thin", color="E2E8F0"),
        right=Side(style="thin", color="E2E8F0"),
        top=Side(style="thin", color="E2E8F0"),
        bottom=Side(style="thin", color="E2E8F0"),
    )
    total_border = Border(
        left=Side(style="thin", color="CBD5E1"),
        right=Side(style="thin", color="CBD5E1"),
        top=Side(style="medium", color="059669"),
        bottom=Side(style="double", color="059669"),
    )

    align_center = Alignment(horizontal="center", vertical="center")
    align_right = Alignment(horizontal="right", vertical="center")

    # ردیف عنوان گزارش
    ws.merge_cells("A1:G1")
    t_cell = ws["A1"]
    t_cell.value = "گزارش ساعات کاری و سوابق کارکرد"
    t_cell.font = title_font
    t_cell.alignment = align_center
    ws.row_dimensions[1].height = 28

    # ردیف توضیحات فیلتر و بازه
    proj_name = project.name if project else "همه پروژه‌ها"
    date_range = f"{start} تا {end}" if (start and end) else "کل سوابق ثبت‌شده"
    now_str = jdatetime.datetime.now().strftime("%Y/%m/%d %H:%M")
    ws.merge_cells("A2:G2")
    s_cell = ws["A2"]
    s_cell.value = f"پروژه: {proj_name} | بازه تاریخی: {date_range} | زمان استخراج: {now_str}"
    s_cell.font = sub_font
    s_cell.alignment = align_center
    ws.row_dimensions[2].height = 20

    # ردیف هدر جدول (ردیف ۴)
    headers = [
        "ردیف",
        "شناسه",
        "تاریخ و ساعت ورود",
        "پروژه",
        "مدت کارکرد (زمان)",
        "کارکرد (ساعت)",
        "توضیحات",
    ]
    ws.append([])  # ردیف ۳ خالی برای فاصله
    ws.append(headers)  # ردیف ۴
    ws.row_dimensions[4].height = 24

    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=4, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = align_center
        cell.border = thin_border

    # درج داده‌های گزارش
    current_row = 5
    for idx, entry in enumerate(daily_report, 1):
        # فرمت تاریخ و ساعت ورود
        entry_time = entry.get("entry_time")
        if entry_time:
            if hasattr(entry_time, "strftime"):
                entry_time_str = entry_time.strftime("%Y-%m-%d %H:%M")
            else:
                entry_time_str = str(entry_time)
        else:
            entry_time_str = "-"

        # محاسبه مدت زمان
        tot_dur = entry.get("total_duration")
        if tot_dur:
            dur_str = str(tot_dur)
            hours_val = round(tot_dur.total_seconds() / 3600, 2)
        else:
            dur_str = "00:00:00"
            hours_val = 0.0

        project_str = entry.get("project__name") or "-"
        desc_str = entry.get("description") or ""

        row_data = [
            idx,
            entry.get("id"),
            entry_time_str,
            project_str,
            dur_str,
            hours_val,
            desc_str,
        ]
        ws.append(row_data)
        ws.row_dimensions[current_row].height = 20

        row_fill = alt_fill if idx % 2 == 0 else white_fill
        for col_idx in range(1, len(row_data) + 1):
            c = ws.cell(row=current_row, column=col_idx)
            c.font = cell_font
            c.fill = row_fill
            c.border = thin_border
            if col_idx in (1, 2, 3, 5):
                c.alignment = align_center
            elif col_idx == 6:
                c.alignment = align_center
                c.number_format = "0.00"
            else:
                c.alignment = align_right

        current_row += 1

    # ردیف جمع‌بندی نهایی
    summary_row = current_row
    ws.merge_cells(start_row=summary_row, start_column=1, end_row=summary_row, end_column=5)
    sum_label = ws.cell(row=summary_row, column=1)
    sum_label.value = "جمع کل ساعات کاری محاسبه‌شده:"
    sum_label.font = total_font
    sum_label.alignment = align_right

    sum_val = ws.cell(row=summary_row, column=6)
    sum_val.value = total_hours
    sum_val.font = total_font
    sum_val.alignment = align_center
    sum_val.number_format = "0.00"

    ws.row_dimensions[summary_row].height = 24
    for col_idx in range(1, 8):
        c = ws.cell(row=summary_row, column=col_idx)
        c.fill = total_fill
        c.border = total_border

    # تنظیم عرض ستون‌ها
    column_widths = {
        "A": 8,   # ردیف
        "B": 10,  # شناسه
        "C": 20,  # تاریخ و ساعت ورود
        "D": 24,  # پروژه
        "E": 18,  # مدت زمان
        "F": 16,  # ساعت
        "G": 35,  # توضیحات
    }
    for col_letter, width in column_widths.items():
        ws.column_dimensions[col_letter].width = width

    # ذخیره در بافر حافظه
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    # ساخت پاسخ دانلودی
    safe_start = start or "all"
    safe_end = end or "all"
    filename = f"report_{safe_start}_{safe_end}.xlsx"
    response = HttpResponse(
        buffer.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    return response
