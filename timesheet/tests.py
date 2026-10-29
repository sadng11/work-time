import jdatetime
from django.test import TestCase, Client
from django.contrib.auth.models import User
from django.urls import reverse

from timesheet.models import Project, TimeEntry
from timesheet.forms import ManualTimeEntryForm, TimeEntryForm


class TimeEntryManualAndCrudTests(TestCase):
    def setUp(self):
        self.user = User.objects.create_user(username='tester', password='password123')
        self.other_user = User.objects.create_user(username='other', password='password123')
        self.project = Project.objects.create(name='پروژه تست', user=self.user, unit_price=400000)
        self.other_project = Project.objects.create(name='پروژه کاربر دیگر', user=self.other_user)

        self.client = Client()
        self.client.login(username='tester', password='password123')

    def test_manual_time_entry_form_validation(self):
        """بررسی اعتبارسنجی فرم ثبت دستی: زمان خروج نباید قبل از زمان ورود باشد"""
        now = jdatetime.datetime.now()
        one_hour_ago = now - jdatetime.timedelta(hours=1)

        # حالت نامعتبر: زمان خروج قبل از ورود
        invalid_data = {
            'project': self.project.id,
            'entry_time': now.strftime('%Y-%m-%d %H:%M'),
            'exit_time': one_hour_ago.strftime('%Y-%m-%d %H:%M'),
            'description': 'تست زمان نامعتبر'
        }
        form = ManualTimeEntryForm(data=invalid_data, user=self.user)
        self.assertFalse(form.is_valid())
        self.assertIn('__all__', form.errors)

        # حالت معتبر: زمان خروج بعد از ورود
        valid_data = {
            'project': self.project.id,
            'entry_time': one_hour_ago.strftime('%Y-%m-%d %H:%M'),
            'exit_time': now.strftime('%Y-%m-%d %H:%M'),
            'description': 'تست زمان معتبر'
        }
        valid_form = ManualTimeEntryForm(data=valid_data, user=self.user)
        self.assertTrue(valid_form.is_valid())

    def test_project_queryset_filtered_by_user(self):
        """پروژه‌های فرم باید فقط شامل پروژه‌های همان کاربر باشند"""
        form = ManualTimeEntryForm(user=self.user)
        project_ids = list(form.fields['project'].queryset.values_list('id', flat=True))
        self.assertIn(self.project.id, project_ids)
        self.assertNotIn(self.other_project.id, project_ids)

    def test_home_post_manual_entry(self):
        """ثبت سابقه دستی از طریق صفحه home"""
        start = jdatetime.datetime(1403, 6, 1, 9, 0)
        end = jdatetime.datetime(1403, 6, 1, 13, 30)

        response = self.client.post(reverse('home'), {
            'action': 'manual_entry',
            'project': self.project.id,
            'entry_time': start.strftime('%Y-%m-%d %H:%M'),
            'exit_time': end.strftime('%Y-%m-%d %H:%M'),
            'description': 'کارکرد دستی گذشته'
        })
        self.assertEqual(response.status_code, 302)

        entry = TimeEntry.objects.filter(user=self.user, description='کارکرد دستی گذشته').first()
        self.assertIsNotNone(entry)
        self.assertEqual(entry.project, self.project)
        self.assertIsNotNone(entry.duration)
        # 4.5 hours = 16200 seconds
        self.assertEqual(entry.duration.seconds, 16200)

    def test_edit_time_entry(self):
        """ویرایش رکورد زمان کاری توسط صاحب رکورد"""
        start = jdatetime.datetime(1403, 6, 2, 10, 0)
        end = jdatetime.datetime(1403, 6, 2, 12, 0)
        entry = TimeEntry.objects.create(
            user=self.user,
            project=self.project,
            entry_time=start,
            exit_time=end,
            description='قبل از ویرایش'
        )

        new_end = jdatetime.datetime(1403, 6, 2, 14, 0)
        edit_url = reverse('edit_time_entry', args=[entry.id])

        # درخواست GET
        get_response = self.client.get(edit_url)
        self.assertEqual(get_response.status_code, 200)

        # درخواست POST برای ویرایش
        post_response = self.client.post(edit_url, {
            'project': self.project.id,
            'entry_time': start.strftime('%Y-%m-%d %H:%M'),
            'exit_time': new_end.strftime('%Y-%m-%d %H:%M'),
            'description': 'بعد از ویرایش'
        })
        self.assertEqual(post_response.status_code, 302)

        entry.refresh_from_db()
        self.assertEqual(entry.description, 'بعد از ویرایش')
        # مدت زمان باید از 2 ساعت به 4 ساعت به‌روزرسانی شده باشد
        self.assertEqual(entry.duration.seconds, 4 * 3600)

    def test_delete_time_entry(self):
        """حذف رکورد زمان کاری با متد POST"""
        entry = TimeEntry.objects.create(
            user=self.user,
            project=self.project,
            entry_time=jdatetime.datetime.now(),
            description='رکورد برای حذف'
        )
        delete_url = reverse('delete_time_entry', args=[entry.id])

        response = self.client.post(delete_url)
        self.assertEqual(response.status_code, 302)
        self.assertFalse(TimeEntry.objects.filter(id=entry.id).exists())

    def test_cannot_edit_or_delete_other_user_entry(self):
        """کاربر نباید بتواند رکورد کاربر دیگری را ویرایش یا حذف کند"""
        other_entry = TimeEntry.objects.create(
            user=self.other_user,
            project=self.other_project,
            entry_time=jdatetime.datetime.now(),
            description='رکورد متعلق به کاربر دیگر'
        )

        edit_url = reverse('edit_time_entry', args=[other_entry.id])
        delete_url = reverse('delete_time_entry', args=[other_entry.id])

        # تلاش برای مشاهده صفحه ویرایش
        edit_response = self.client.get(edit_url)
        self.assertEqual(edit_response.status_code, 404)

        # تلاش برای حذف
        delete_response = self.client.post(delete_url)
        self.assertEqual(delete_response.status_code, 404)

        # رکورد هنوز در دیتابیس موجود است
        self.assertTrue(TimeEntry.objects.filter(id=other_entry.id).exists())

    def test_factor_page_back_to_report_link(self):
        """دکمه بازگشت در صفحه فاکتور باید به صفحه گزارش با فیلترهای مشخص شده لینک داشته باشد"""
        start_str = '1405-07-01'
        end_str = '1405-08-05'
        factor_url = reverse('factor', args=[start_str, end_str, self.project.id])

        response = self.client.get(factor_url)
        self.assertEqual(response.status_code, 200)

        expected_report_url = f"{reverse('report')}?start={start_str}&end={end_str}&project_id={self.project.id}"
        self.assertContains(response, expected_report_url)

    def test_report_view_without_project_filter(self):
        """صفحه گزارش بدون انتخاب پروژه نباید با خطای VariableDoesNotExist روبرو شود"""
        entry = TimeEntry.objects.create(
            user=self.user,
            project=self.project,
            entry_time=jdatetime.datetime(1405, 8, 2, 10, 0),
            exit_time=jdatetime.datetime(1405, 8, 2, 12, 0),
            description='تست گزارش ماه'
        )
        report_url = f"{reverse('report')}?start=1405-08-01&end=1405-08-05"
        response = self.client.get(report_url)
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'همه پروژه‌ها')
        self.assertContains(response, self.project.name)

    def test_factor_page_decimal_hours_and_fee(self):
        """فاکتور نباید ساعت را به عدد صحیح گرد کند و محاسبات باید با اعشار دقیق باشد"""
        start = jdatetime.datetime(1405, 8, 1, 8, 0)
        end = jdatetime.datetime(1405, 8, 1, 13, 12)
        TimeEntry.objects.create(
            user=self.user,
            project=self.project,
            entry_time=start,
            exit_time=end,
            description='کارکرد با اعشار'
        )
        factor_url = reverse('factor', args=['1405-08-01', '1405-08-05', self.project.id])
        response = self.client.get(factor_url)
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, '5.20')
        self.assertContains(response, '2,080,000')

    def test_report_excel_export_link_in_html(self):
        """صفحه گزارش باید دکمه خروجی اکسل با پارامتر export=excel را نمایش دهد"""
        TimeEntry.objects.create(
            user=self.user,
            project=self.project,
            entry_time=jdatetime.datetime(1405, 8, 1, 9, 0),
            exit_time=jdatetime.datetime(1405, 8, 1, 11, 0),
            description='تست لینک اکسل'
        )
        url = f"{reverse('report')}?start=1405-08-01&end=1405-08-05&project_id={self.project.id}"
        response = self.client.get(url)
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'export=excel')
        self.assertContains(response, 'خروجی اکسل')

    def test_report_excel_export_file(self):
        """درخواست خروجی اکسل باید فایل معتبر با فرمت و هدر مناسب تولید کند"""
        import io
        import openpyxl

        TimeEntry.objects.create(
            user=self.user,
            project=self.project,
            entry_time=jdatetime.datetime(1405, 8, 1, 9, 0),
            exit_time=jdatetime.datetime(1405, 8, 1, 12, 0),
            description='کارکرد برای اکسل'
        )
        url = f"{reverse('report')}?start=1405-08-01&end=1405-08-05&project_id={self.project.id}&export=excel"
        response = self.client.get(url)
        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response['Content-Type'],
            'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        self.assertIn('attachment; filename="report_1405-08-01_1405-08-05.xlsx"', response['Content-Disposition'])

        # اعتبارسنجی ساختار فایل اکسل دریافتی
        wb = openpyxl.load_workbook(io.BytesIO(response.content))
        self.assertIn('گزارش کارکرد', wb.sheetnames)
        ws = wb['گزارش کارکرد']
        self.assertTrue(ws.views.sheetView[0].rightToLeft)
        # بررسی وجود عنوان گزارش در سلول A1
        self.assertEqual(ws['A1'].value, 'گزارش ساعات کاری و سوابق کارکرد')
        # بررسی وجود ستون مبلغ و مقادیر محاسبه‌شده
        self.assertEqual(ws.cell(row=4, column=7).value, 'مبلغ (تومان)')
        # ۳ ساعت با نرخ ۴۰۰,۰۰۰ تومان = ۱,۲۰۰,۰۰۰ تومان
        self.assertEqual(ws.cell(row=5, column=7).value, 1200000)
        # سلول جمع کل مبلغ در ردیف پایانی
        self.assertEqual(ws.cell(row=6, column=7).value, 1200000)

    def test_report_page_fee_column_and_total(self):
        """صفحه گزارش باید ستون مبلغ و جمع کل مبلغ به تومان را نمایش دهد"""
        start = jdatetime.datetime(1405, 8, 1, 8, 0)
        end = jdatetime.datetime(1405, 8, 1, 10, 0)
        # رکورد اول برای پروژه دارای نرخ ۴۰۰,۰۰۰ تومان (۲ ساعت = ۸۰۰,۰۰۰ تومان)
        TimeEntry.objects.create(
            user=self.user,
            project=self.project,
            entry_time=start,
            exit_time=end,
            description='کار با نرخ مشخص'
        )
        # پروژه بدون نرخ پایه (unit_price=None)
        no_price_project = Project.objects.create(name='پروژه بدون نرخ', user=self.user, unit_price=None)
        TimeEntry.objects.create(
            user=self.user,
            project=no_price_project,
            entry_time=start + jdatetime.timedelta(hours=3),
            exit_time=end + jdatetime.timedelta(hours=3),
            description='کار بدون نرخ'
        )

        url = f"{reverse('report')}?start=1405-08-01&end=1405-08-05"
        response = self.client.get(url)
        self.assertEqual(response.status_code, 200)
        # وجود هدر ستون مبلغ
        self.assertContains(response, 'مبلغ (تومان)')
        # مبلغ رکورد اول
        self.assertContains(response, '800,000')
        # مبلغ رکورد دوم بدون نرخ باید ۰ تومان باشد
        self.assertContains(response, '0 <span class="small text-muted fw-normal">تومان</span>')
        # جمع کل مبلغ (۸۰۰,۰۰۰ تومان) در پاورقی و کارت KPI
        self.assertContains(response, 'جمع کل کارکرد و مبلغ:')
        self.assertContains(response, '800,000 تومان')

    def test_home_page_fee_column_and_total(self):
        """صفحه اصلی باید ستون مبلغ و سطر جمع‌بندی را در جدول نمایش دهد"""
        start = jdatetime.datetime.now() - jdatetime.timedelta(hours=2)
        end = jdatetime.datetime.now()
        TimeEntry.objects.create(
            user=self.user,
            project=self.project,
            entry_time=start,
            exit_time=end,
            description='کار اخیر در صفحه اصلی'
        )
        response = self.client.get(reverse('home'))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'مبلغ (تومان)')
        self.assertContains(response, 'جمع کل سوابق اخیر:')
        self.assertContains(response, 'تومان')

    def test_start_timer_creates_active_entry(self):
        """شروع تایمر باید یک رکورد با exit_time خالی ایجاد کند"""
        url = reverse('start_timer')
        response = self.client.post(url, {
            'project': self.project.id,
            'description': 'تست شروع تایمر زنده',
        })
        self.assertRedirects(response, reverse('home'))
        active = TimeEntry.objects.filter(user=self.user, exit_time__isnull=True).first()
        self.assertIsNotNone(active)
        self.assertEqual(active.project, self.project)
        self.assertEqual(active.description, 'تست شروع تایمر زنده')
        self.assertEqual(active.entry_time.microsecond, 0)

    def test_start_timer_auto_stops_previous_active_timer(self):
        """شروع تایمر روی پروژه جدید باید تایمر فعال قبلی را به طور خودکار متوقف و ثبت کند"""
        other_project = Project.objects.create(name='پروژه دوم', user=self.user, unit_price=250_000)
        # شروع تایمر اول
        self.client.post(reverse('start_timer'), {'project': self.project.id, 'description': 'کار اول'})
        first_entry = TimeEntry.objects.filter(user=self.user, project=self.project, exit_time__isnull=True).first()
        self.assertIsNotNone(first_entry)

        # شروع تایمر دوم
        self.client.post(reverse('start_timer'), {'project': other_project.id, 'description': 'کار دوم'})

        first_entry.refresh_from_db()
        self.assertIsNotNone(first_entry.exit_time)
        self.assertIsNotNone(first_entry.duration)

        second_entry = TimeEntry.objects.filter(user=self.user, project=other_project, exit_time__isnull=True).first()
        self.assertIsNotNone(second_entry)

    def test_stop_timer_sets_exit_time_and_duration(self):
        """توقف تایمر باید exit_time را تنظیم کرده و مدت زمان را محاسبه کند"""
        # ایجاد تایمر باز
        entry = TimeEntry.objects.create(
            user=self.user,
            project=self.project,
            entry_time=jdatetime.datetime.now() - jdatetime.timedelta(minutes=30),
            description='کار در حال اجرا'
        )
        url = reverse('stop_timer')
        response = self.client.post(url)
        self.assertRedirects(response, reverse('home'))

        entry.refresh_from_db()
        self.assertIsNotNone(entry.exit_time)
        self.assertEqual(entry.exit_time.microsecond, 0)
        self.assertIsNotNone(entry.duration)
        self.assertGreaterEqual(entry.duration.total_seconds(), 1700)

    def test_timer_datetime_has_no_microseconds_in_db_and_templates(self):
        """زمان‌های ثبت شده نباید دارای اعشار/میکروثانیه بعد از ثانیه باشند و در فرمت YYYY-MM-DD HH:MM:SS نمایش یابند"""
        now = jdatetime.datetime.now()
        entry = TimeEntry.objects.create(
            user=self.user,
            project=self.project,
            entry_time=now,
            exit_time=now + jdatetime.timedelta(hours=1),
        )
        self.assertEqual(entry.entry_time.microsecond, 0)
        self.assertEqual(entry.exit_time.microsecond, 0)

        response = self.client.get(reverse('home'))
        self.assertEqual(response.status_code, 200)
        formatted_entry = entry.entry_time.strftime('%Y-%m-%d %H:%M:%S')
        formatted_exit = entry.exit_time.strftime('%Y-%m-%d %H:%M:%S')
        self.assertContains(response, formatted_entry)
        self.assertContains(response, formatted_exit)

    def test_home_page_displays_active_live_timer(self):
        """صفحه اصلی در صورت وجود تایمر باز باید ویجت فعال و زمان را نمایش دهد"""
        entry = TimeEntry.objects.create(
            user=self.user,
            project=self.project,
            entry_time=jdatetime.datetime.now(),
            description='تسک در حال انجام'
        )
        response = self.client.get(reverse('home'))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'is-active')
        self.assertContains(response, 'در حال ضبط کارکرد')
        self.assertContains(response, 'توقف و ثبت کارکرد')
        self.assertContains(response, self.project.name)

    def test_analytics_view_status_and_content(self):
        """صفحه تحلیل و آمار باید با وضعیت ۲۰۰ باز شده و نمودارها و شاخص‌ها را لود کند"""
        # ایجاد یک رکورد برای این ماه
        TimeEntry.objects.create(
            user=self.user,
            project=self.project,
            entry_time=jdatetime.datetime.now() - jdatetime.timedelta(hours=2),
            exit_time=jdatetime.datetime.now(),
            description='کار برای تحلیل'
        )
        response = self.client.get(reverse('analytics'))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'تحلیل عملکرد و آمار کارکرد')
        self.assertContains(response, 'weeklyChart')
        self.assertContains(response, 'projectChart')
        self.assertContains(response, 'monthlyTrendChart')
        self.assertContains(response, 'chart.min.js')
        self.assertContains(response, 'analyticsData')

    def test_analytics_context_with_no_entries(self):
        """تابع محاسبات آماری برای کاربر جدید بدون رکورد نباید دچار خطای تقسیم بر صفر شود"""
        from timesheet.analytics import get_analytics_context
        # ساخت کاربر بدون رکورد
        new_user = User.objects.create_user(username='empty_user', password='password123')
        data = get_analytics_context(new_user)
        self.assertEqual(data['kpi']['total_month_hours'], 0.0)
        self.assertEqual(data['kpi']['daily_avg_hours'], 0.0)
        self.assertEqual(data['kpi']['active_work_days'], 0)
        self.assertEqual(data['kpi']['total_month_income'], 0)
        self.assertEqual(len(data['weekly_chart']['hours']), 7)
        self.assertEqual(len(data['monthly_chart']['hours']), 6)




